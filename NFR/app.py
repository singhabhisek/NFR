import json
from collections import defaultdict
import logging
from logging.handlers import TimedRotatingFileHandler

import openpyxl
from flask import Flask, render_template, request, jsonify, redirect, url_for, flash, session, make_response
import sqlite3
import os
from openpyxl import load_workbook
from datetime import datetime, timedelta
import time

app = Flask(__name__)
app.config['UPLOAD_FOLDER'] = './uploads'
app.config['DOWNLOAD_FOLDER'] = './downloads'
app.config['LOGS_FOLDER'] = './logs'
app.secret_key = 'supersecretkey'

# Ensure the upload folder exists
os.makedirs(app.config['UPLOAD_FOLDER'], exist_ok=True)

# Set up logging
log_file = os.path.join(app.config['LOGS_FOLDER'], 'app.log')
handler = TimedRotatingFileHandler(log_file, when='midnight', interval=1, backupCount=7)
handler.setFormatter(logging.Formatter('%(asctime)s %(levelname)s: [%(func)s] %(message)s'))
handler.setLevel(logging.INFO)
app.logger.addHandler(handler)


# SQLite connection function
def get_db_connection():
    conn = sqlite3.connect('nfr_repository.db')
    conn.row_factory = sqlite3.Row
    return conn


# Route to serve the static file template.xlsx
@app.route('/downloads/template.xlsx')
def download_template():
    return app.send_static_file('downloads/template.xlsx')

# Home route
@app.route('/')
def index():
    conn = get_db_connection()
    apps = conn.execute('SELECT DISTINCT applicationName FROM NFRDetails').fetchall()
    conn.close()

    return render_template('index.html', apps=apps, role='admin')

# Route to upload the file
@app.route('/upload', methods=['GET', 'POST'])
def upload_file():
    error_log = []  # Initialize error_log
    summary = None  # Initialize summary
    if request.method == 'GET':
        pass
    if request.method == 'POST':
        file = request.files['file']
        if file and file.filename.endswith('.xlsx'):
            filepath = os.path.join(app.config['UPLOAD_FOLDER'], file.filename)
            file.save(filepath)
            try:
                error_log, summary = process_excel_file(filepath)
                if summary is None:
                    flash('File was not processed due to errors.', 'danger')
                else:
                    flash('File successfully uploaded and processed.', 'success')
            except Exception as e:
                app.logger.error(f'Error processing file: {e}', extra={'func': 'UPLOAD'})
                flash(f'Error processing file: {e}', 'danger')
            return render_template('upload.html', logs=error_log, summary=summary)
    return render_template('upload.html', logs=[], summary=summary)


# Function to process the uploaded Excel file
def process_excel_file(filepath):
    error_log = []
    total_records = 0
    inserted_records = 0
    updated_records = 0
    passed_records = 0
    failed_records = 0

    # Define mandatory columns
    mandatory_columns = ['ApplicationName', 'ReleaseID', 'BusinessScenario', 'TransactionName', 'SLA', 'TPS',
                         'Comments']

    try:
        wb = load_workbook(filename=filepath)
        ws = wb.active
    except Exception as load_error:
        error_log.append(f"Error loading Excel file: {load_error}")
        app.logger.error(f"Error loading Excel file: {load_error}", extra={'func': 'PROCESS_EXCEL_FILE'})
        return error_log, None  # Return error_log and summary as None immediately if loading the file fails

    # Extract headers from the first row
    headers = [cell.value for cell in ws[1]]

    # Check if mandatory columns are present
    missing_columns = [col for col in mandatory_columns if col not in headers]
    if missing_columns:
        error_message = f"Missing mandatory columns: {', '.join(missing_columns)}"
        error_log.append(error_message)
        app.logger.error(error_message, extra={'func': 'PROCESS_EXCEL_FILE'})
        raise Exception(error_message)  # Abort processing if mandatory columns are missing

    # Map headers to their indexes
    header_index_map = {header: index for index, header in enumerate(headers)}

    current_user = session.get('user', 'defaultUser')  # Retrieve current user
    current_time = "2024-06-06 10:10:10" #datetime.now().strftime('%Y-%m-%d %H:%M:%S')

    with get_db_connection() as conn:
        for row_number, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            # Extract data using header indices
            try:
                applicationName = row[header_index_map['ApplicationName']]
                releaseID = row[header_index_map['ReleaseID']]
                businessScenario = row[header_index_map['BusinessScenario']]
                transactionName = row[header_index_map['TransactionName']]
                SLA = row[header_index_map['SLA']]
                TPS = row[header_index_map['TPS']]
                comments = row[header_index_map['Comments']]
            except IndexError as e:
                failed_records += 1
                error_message = f"Error processing row {row_number}: {e}"
                error_log.append(error_message)
                app.logger.error(error_message, extra={'func': 'PROCESS_EXCEL_FILE'})
                continue

            total_records += 1
            # Validate SLA and TPS
            if SLA is None or TPS is None:
                failed_records += 1
                error_message = f"Error processing row {row_number}: SLA and TPS must have values"
                error_log.append(error_message)
                app.logger.error(error_message, extra={'func': 'PROCESS_EXCEL_FILE'})
                continue

            try:
                SLA = float(SLA)
                TPS = float(TPS)
            except ValueError:
                failed_records += 1
                error_message = f"Error processing row {row_number}: SLA and TPS must be valid numbers"
                error_log.append(error_message)
                app.logger.error(error_message, extra={'func': 'PROCESS_EXCEL_FILE'})
                continue

            try:
                # Check if the record exists
                cursor = conn.execute('''
                    SELECT * FROM NFRDetails
                    WHERE applicationName = ? AND releaseID = ? AND  transactionName = ?
                ''', (applicationName, releaseID, transactionName))

                existing_record = cursor.fetchone()
                if existing_record:
                    # Update the record
                    conn.execute('''
                        UPDATE NFRDetails
                        SET businessScenario = ?, SLA = ?, TPS = ?, comments = ?, modifiedBy = ?, modified_date = ?
                        WHERE applicationName = ? AND releaseID = ? AND transactionName = ?
                    ''', (businessScenario, SLA, TPS, comments, current_user, current_time, applicationName, releaseID, transactionName))
                    updated_records += 1
                    # print(businessScenario, transactionName, SLA, TPS, comments, current_user, current_time, applicationName, releaseID)
                else:
                    # Insert the record
                    conn.execute('''
                        INSERT INTO NFRDetails (applicationName, releaseID, businessScenario, transactionName, SLA, TPS, comments, createdBy, modifiedBy, created_date, modified_date)
                        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    ''', (applicationName, releaseID, businessScenario, transactionName, SLA, TPS, comments, current_user, current_user, current_time, current_time))
                    inserted_records += 1
                passed_records += 1
            except Exception as e:
                failed_records += 1
                error_message = f"Error processing row {row_number}: {e}"
                error_log.append(error_message)
                app.logger.error(error_message, extra={'func': 'PROCESS_EXCEL_FILE'})

        conn.commit()

    summary = {
        'total': total_records,
        'inserted': inserted_records,
        'updated': updated_records,
        'passed': passed_records,
        'failed': failed_records
    }
    # print(summary)
    return error_log, summary

# Fetch releases based on application name
@app.route('/get_releases', methods=['POST'])
def get_releases():
    try:
        applicationName = request.form['applicationName']
        conn = get_db_connection()
        releases = conn.execute('SELECT DISTINCT releaseID FROM NFRDetails WHERE applicationName = ?', (applicationName,)).fetchall()
        conn.close()
        return jsonify([release['releaseID'] for release in releases]), 200
    except Exception as e:
        return jsonify(success=False, message=str(e)), 500

# Suggest transactions based on term and application name
@app.route('/suggest_transactions', methods=['POST'])
def suggest_transactions():
    try:
        term = request.form['term']
        applicationName = request.form['applicationName']
        conn = get_db_connection()
        transactions = conn.execute('SELECT DISTINCT transactionName FROM NFRDetails WHERE applicationName = ? AND transactionName LIKE ?', (applicationName, f'%{term}%')).fetchall()
        conn.close()
        return jsonify([txn['transactionName'] for txn in transactions]), 200
    except Exception as e:
        return jsonify(success=False, message=str(e)), 500
#
# # Search records based on various parameters
# @app.route('/search_records', methods=['POST'])
# def search_records():
#     try:
#         applicationName = request.form['application_name']
#         releaseID = request.form.get('release_id')
#         transactionName = request.form.get('transaction_name')
#         print('asasa')
#         print('aaa' + applicationName, releaseID, transactionName)
#         conn = get_db_connection()
#
#         query = 'SELECT * FROM NFRDetails WHERE applicationName = ?'
#         params = [applicationName]
#
#         if releaseID:
#             query += ' AND releaseID = ?'
#             params.append(releaseID)
#         if transactionName:
#             if '*' in transactionName:
#                 transactionName = transactionName.replace('*', '%')
#             else:
#                 transactionName = f'%{transactionName}%'
#             query += ' AND transactionName LIKE ?'
#             params.append(transactionName)
#
#         records = conn.execute(query, params).fetchall()
#         conn.close()
#
#         return jsonify([dict(record) for record in records]), 200
#     except Exception as e:
#         return jsonify(success=False, message=str(e)), 500


# Search records based on various parameters
@app.route('/search_records', methods=['POST'])
def search_records():
    # Simulate processing delay
    time.sleep(1)  # Delay of 1 second
    try:
        applicationName = request.form['application_name']
        releaseID = request.form.get('release_id')
        transactionName = request.form.get('transaction_name')

        conn = get_db_connection()

        query = 'SELECT * FROM NFRDetails WHERE applicationName = ?'
        params = [applicationName]

        if releaseID:
            query += ' AND releaseID = ?'
            params.append(releaseID)
        if transactionName:
            if '*' in transactionName:
                transactionName = transactionName.replace('*', '%')
            else:
                transactionName = f'%{transactionName}%'
            query += ' AND transactionName LIKE ?'
            params.append(transactionName)

        records = conn.execute(query, params).fetchall()
        conn.close()

        return jsonify([dict(record) for record in records]), 200
    except Exception as e:
        return jsonify(success=False, message=str(e)), 500


# Get a single record by ID
@app.route('/get_record', methods=['POST'])
def get_record():
    try:
        record_id = request.form.get('Id')
        if not record_id:
            return jsonify(success=False, message="Record ID is missing."), 400
        conn = get_db_connection()
        record = conn.execute('SELECT * FROM NFRDetails WHERE Id = ?', (record_id,)).fetchone()
        conn.close()
        if not record:
            return jsonify(success=False, message="Record not found."), 404
        return jsonify(dict(record)), 200
    except Exception as e:
        return jsonify(success=False, message=str(e)), 500

# Update a record
@app.route('/update_record', methods=['POST'])
def update_record():
    try:
        data = request.form.to_dict()

        # Validate SLA and TPS
        sla = data.get('SLA')
        tps = data.get('TPS')

        try:
            sla_value = float(sla)
            tps_value = float(tps)
        except ValueError:
            return jsonify(success=False, message="SLA and TPS must be valid numbers."), 400

        if sla_value <= 0 or tps_value <= 0:
            return jsonify(success=False, message="SLA and TPS must be positive numbers."), 400

        conn = get_db_connection()
        conn.execute(
            'UPDATE NFRDetails SET businessScenario = ?, transactionName = ?, SLA = ?, TPS = ? WHERE Id = ?',
            (data['businessScenario'], data['transactionName'], data['SLA'], data['TPS'], data['Id'])
        )
        conn.commit()
        conn.close()
        return jsonify(success=True, message="Record updated successfully."), 200
    except Exception as e:
        return jsonify(success=False, message=str(e)), 500

# Delete a record
@app.route('/delete_record', methods=['POST'])
def delete_record():
    try:
        record_id = request.form['Id']
        conn = get_db_connection()
        conn.execute('DELETE FROM NFRDetails WHERE Id = ?', (record_id,))
        conn.commit()
        conn.close()
        return jsonify(success=True, message="Record deleted successfully."), 200
    except Exception as e:
        return jsonify(success=False, message=str(e)), 500


# Route to compare records
@app.route('/compare')
def compare():
    conn = get_db_connection()
    apps = conn.execute('SELECT DISTINCT applicationName FROM NFRDetails').fetchall()
    conn.close()
    return render_template('compare.html', apps=apps)


# @app.route('/compare_records', methods=['POST'])
# def compare_records():
#     try:
#         applicationName = request.form['applicationName']
#         releaseID1 = request.form['releaseID1']
#         releaseID2 = request.form['releaseID2']
#         releaseID3 = request.form['releaseID3']
#         transactionName = request.form['transactionName']
#
#         if '*' in transactionName:
#             transactionName = transactionName.replace('*', '%')
#         else:
#             transactionName = f'{transactionName}'
#
#         selected_releases = [(releaseID1, 'Release1'), (releaseID2, 'Release2'), (releaseID3, 'Release3')]
#         columns = []
#         query_cols = []
#         for release, alias in selected_releases:
#             if release:
#                 columns.append(f'{release} SLA')
#                 columns.append(f'{release} TPS')
#                 query_cols.append(f'SUM(CASE WHEN releaseID = ? THEN SLA END) AS "{release} SLA"')
#                 query_cols.append(f'SUM(CASE WHEN releaseID = ? THEN TPS END) AS "{release} TPS"')
#
#         query = f"""
#         SELECT transactionName, {", ".join(query_cols)}
#         FROM NFRDetails
#         WHERE applicationName = ?
#         AND transactionName LIKE ?
#         GROUP BY transactionName
#         """
#
#         # Ensure the parameter order matches the query
#         params = []
#         for release, _ in selected_releases:
#             if release:
#                 params.extend([release, release])  # Extend params list with release twice (for SLA and TPS)
#         params.append(applicationName)
#         params.append(f'%{transactionName}%')
#
#         conn = get_db_connection()
#         cursor = conn.execute(query, params)
#         results = cursor.fetchall()
#         conn.close()
#
#         # Convert results to a dictionary
#         records = [{**dict(row)} for row in results]
#         print(records)
#         # Return JSON response with columns and records
#         return jsonify(columns=['Transaction Name'] + columns, records=records), 200
#
#     except Exception as e:
#         print("Exception occurred:", str(e))
#         return jsonify(success=False, message=str(e)), 500

@app.route('/compare_records', methods=['POST'])
def compare_records():
    try:
        applicationName = request.form['applicationName']
        releaseID1 = request.form['releaseID1']
        releaseID2 = request.form['releaseID2']
        releaseID3 = request.form['releaseID3']
        transactionName = request.form['transactionName']
        # Check if show_all_rows parameter is present in the form data
        if 'show_all_rows' in request.form:
            show_all_rows = request.form['show_all_rows'].lower() == 'true'  # Convert to boolean
        else:
            show_all_rows = False  # Default to False if parameter is not provided

        if '*' in transactionName:
            transactionName = transactionName.replace('*', '%')
        else:
            transactionName = f'{transactionName}'

        selected_releases = [(releaseID1, 'Release1'), (releaseID2, 'Release2'), (releaseID3, 'Release3')]
        columns = []
        query_cols = []
        for release, alias in selected_releases:
            if release:
                columns.append(f'{release} SLA')
                columns.append(f'{release} TPS')
                query_cols.append(f'SUM(CASE WHEN releaseID = ? THEN SLA END) AS "{release} SLA"')
                query_cols.append(f'SUM(CASE WHEN releaseID = ? THEN TPS END) AS "{release} TPS"')

        query = f"""
        SELECT transactionName, {", ".join(query_cols)}
        FROM NFRDetails
        WHERE applicationName = ?
        AND transactionName LIKE ?
        GROUP BY transactionName
        """

        # Ensure the parameter order matches the query
        params = []
        for release, _ in selected_releases:
            if release:
                params.extend([release, release])  # Extend params list with release twice (for SLA and TPS)
        params.append(applicationName)
        params.append(f'%{transactionName}%')

        conn = get_db_connection()
        cursor = conn.execute(query, params)
        results = cursor.fetchall()
        conn.close()

        # Convert results to a dictionary
        records = [{**dict(row)} for row in results]

        # Filter records based on show_all_rows flag
        if not show_all_rows:
            records = [record for record in records if any(record.get(column) is not None for column in columns)]

        # Return JSON response with columns and filtered records
        return jsonify(columns=['Transaction Name'] + columns, records=records), 200

    except Exception as e:
        print("Exception occurred:", str(e))
        return jsonify(success=False, message=str(e)), 500

# # Route for the discrepancy page
# @app.route('/discrepancy', methods=['GET', 'POST'])
# def handle_discrepancy():
#     if request.method == 'POST':
#         application_name = request.form.get('application_name')
#         release_id = request.form.get('release_id')
#         #     WHEN COALESCE(TotalSyncSLA, 0) + COALESCE(MaxAsyncSLA, 0) = 0 THEN 'NA'
#         #     ELSE CASE WHEN SLA > COALESCE(TotalSyncSLA, 0) + COALESCE(MaxAsyncSLA, 0) THEN 'Higher' ELSE 'Lower' END
#         # END AS Compare
#         query = """
#             WITH NFRDetailDepend AS (
#                 SELECT
#                     t.applicationName, t.transactionName, t.releaseID, t.SLA,
#                     d.backendCall, d.callType
#                 FROM NFRDetails t
#                 LEFT JOIN  NFROperationDependency d ON t.transactionName = d.transactionName
#                 WHERE t.applicationName = ? AND t.releaseID = ?
#             )
#             SELECT
#                 ApplicationName, releaseID, transactionName, SLA,
#                 COALESCE(TotalSyncSLA, 0) + COALESCE(MaxAsyncSLA, 0) AS TotalBackendCallDuration,
#                 GROUP_CONCAT(backendCall) AS backendCallList,
#                 CASE
#
#                     WHEN COALESCE(TotalSyncSLA, 0) + COALESCE(MaxAsyncSLA, 0) = 0 THEN 'NA'
#                     ELSE CASE WHEN SLA > COALESCE(TotalSyncSLA, 0) + COALESCE(MaxAsyncSLA, 0) THEN 'Higher'
#                               WHEN SLA < COALESCE(TotalSyncSLA, 0) + COALESCE(MaxAsyncSLA, 0) THEN 'Lower'
#                               ELSE 'Equal'
#                           END
#                 END AS Compare
#             FROM (
#                 SELECT
#                     ApplicationName, transactionName, releaseID, SLA, backendCall,
#                     SUM(CASE WHEN CallType = 'Sync' THEN SLAComparison ELSE 0 END) OVER (PARTITION BY transactionName) AS TotalSyncSLA,
#                     MAX(CASE WHEN CallType = 'Async' THEN SLAComparison ELSE 0 END) OVER (PARTITION BY transactionName) AS MaxAsyncSLA
#                 FROM (
#                     SELECT
#                         ApplicationName, transactionName, backendCall, CallType, SLA, releaseID,
#                         CASE
#                             WHEN CallType = 'Async' THEN (SELECT MAX(SLA) FROM NFRDetails WHERE transactionName = t.backendCall AND t.CallType = 'Async')
#                             WHEN CallType = 'Sync' THEN (SELECT SUM(SLA) FROM NFRDetails WHERE transactionName = t.backendCall AND t.CallType = 'Sync')
#                             ELSE 0
#                         END AS SLAComparison
#                     FROM NFRDetailDepend t
#                 ) AS x
#                 WHERE ApplicationName = ? AND releaseID = ?
#             ) AS p
#             GROUP BY ApplicationName, releaseID, transactionName, SLA, Compare;
#         """
#
#         params = (application_name, release_id, application_name, release_id)
#
#         conn = get_db_connection()
#         cur = conn.cursor()
#         cur.execute(query, params)
#         rows = cur.fetchall()
#         conn.close()
#
#         data = []
#
#         for row in rows:
#             # Use the correct key for backend call list
#             if row['backendCallList'] is not None:
#                 concatenated_calls = f"<b>Backend Duration (sec)</b>: {row['TotalBackendCallDuration']}<br><b>Backend Calls:</b> <br>" + \
#                                      row['backendCallList'].replace(',', ',<br>')
#             else:
#                 concatenated_calls = f"<b>Backend Duration (sec)</b>: {row['TotalBackendCallDuration']}<br><b>Backend Calls:</b> None"
#
#             comparison_icon = ''
#             comparison_text = ''
#
#             if row['Compare'] == 'Higher':
#                 comparison_icon = '<i class="fas fa-arrow-up" title="Higher"></i>'
#                 comparison_text = 'Higher'
#             elif row['Compare'] == 'Lower':
#                 comparison_icon = '<i class="fas fa-arrow-down" title="Lower"></i>'
#                 comparison_text = 'Lower'
#             elif row['Compare'] == 'Equal':
#                 comparison_icon = '<i class="fas fa-exchange-alt" title="Equal"></i>'
#                 comparison_text = 'Equal'
#             # print(concatenated_calls)
#             data.append({
#                 'ApplicationName': row['ApplicationName'],
#                 'releaseID': row['releaseID'],
#                 'transactionName': row['transactionName'],
#                 'SLA': row['SLA'],
#                 'backendCall': concatenated_calls,
#                 'Compare': row['Compare']
#             })
#         # print(data)
#         return jsonify(data)
#
#     else:
#         # Handle GET request logic here (if needed)
#         conn = get_db_connection()
#         apps = conn.execute('SELECT DISTINCT applicationName FROM NFRDetails').fetchall()
#         conn.close()
#
#         searches_cookie = request.cookies.get('recent_searches')
#         if searches_cookie:
#             recent_searches = json.loads(searches_cookie)
#         else:
#             recent_searches = []
#
#         return render_template('discrepancy.html', apps=apps, recent_searches=recent_searches)
#         # return render_template('discrepancy.html', apps=apps)

# # Route for the discrepancy page
# @app.route('/discrepancy', methods=['GET', 'POST'])
# def handle_discrepancy():
#     if request.method == 'POST':
#         application_name = request.form.get('application_name')
#         release_id = request.form.get('release_id')
#         transaction_name = request.form.get('transaction_name')  # Get the transaction name from the form
#
#         # Define the base query with placeholders for filters
#         query = """
#             WITH NFRDetailDepend AS (
#                 SELECT
#                     t.applicationName, t.transactionName, t.releaseID, t.SLA,
#                     d.backendCall, d.callType
#                 FROM NFRDetails t
#                 LEFT JOIN NFROperationDependency d ON t.transactionName = d.transactionName
#                 WHERE t.applicationName = ? AND (
#                         t.releaseID = ? OR
#                         t.releaseID = (SELECT MAX(releaseID) FROM NFRDetails WHERE applicationName = t.applicationName and transactionName = t.transactionName)
#                     )
#         """
#
#         params = [application_name, release_id]
#
#         # Add transactionName filter if provided
#         if transaction_name:
#             query += " AND t.transactionName like ?"
#             # params.append(transaction_name)
#             params.append(f'%{transaction_name}%')
#         query += """
#             )
#             SELECT
#                 ApplicationName, releaseID, transactionName, SLA,
#                 COALESCE(TotalSyncSLA, 0) + COALESCE(MaxAsyncSLA, 0) AS TotalBackendCallDuration,
#                 GROUP_CONCAT(backendCall) AS backendCallList,
#                 CASE
#                     WHEN COALESCE(TotalSyncSLA, 0) + COALESCE(MaxAsyncSLA, 0) = 0 THEN 'NA'
#                     ELSE CASE WHEN SLA > COALESCE(TotalSyncSLA, 0) + COALESCE(MaxAsyncSLA, 0) THEN 'Higher'
#                               WHEN SLA < COALESCE(TotalSyncSLA, 0) + COALESCE(MaxAsyncSLA, 0) THEN 'Lower'
#                               ELSE 'Equal'
#                           END
#                 END AS Compare
#             FROM (
#                 SELECT
#                     ApplicationName, transactionName, releaseID, SLA, backendCall,
#                     SUM(CASE WHEN CallType = 'Sync' THEN SLAComparison ELSE 0 END) OVER (PARTITION BY transactionName) AS TotalSyncSLA,
#                     MAX(CASE WHEN CallType = 'Async' THEN SLAComparison ELSE 0 END) OVER (PARTITION BY transactionName) AS MaxAsyncSLA
#                 FROM (
#                     SELECT
#                         ApplicationName, transactionName, backendCall, CallType, SLA, releaseID,
#                         CASE
#                             WHEN CallType = 'Async' THEN (SELECT MAX(SLA) FROM NFRDetails WHERE transactionName = t.backendCall AND t.CallType = 'Async')
#                             WHEN CallType = 'Sync' THEN (SELECT SUM(SLA) FROM NFRDetails WHERE transactionName = t.backendCall AND t.CallType = 'Sync')
#                             ELSE 0
#                         END AS SLAComparison
#                     FROM NFRDetailDepend t
#                 ) AS x
#                 WHERE ApplicationName = ? AND releaseID = ?
#         """
#
#         params += [application_name, release_id]
#
#         # Add transactionName filter if provided
#         if transaction_name:
#             query += " AND transactionName like ?"
#             # params.append(transaction_name)
#             params.append(f'%{transaction_name}%')
#         query += """
#             ) AS p
#             GROUP BY ApplicationName, releaseID, transactionName, SLA, Compare;
#         """
#
#         print(query)
#         print(params)
#         # Execute the query
#         conn = get_db_connection()
#         cur = conn.cursor()
#         cur.execute(query, params)
#         rows = cur.fetchall()
#         conn.close()
#
#         data = []
#
#         for row in rows:
#             # Use the correct key for backend call list
#             if row['backendCallList'] is not None:
#                 concatenated_calls = f"<b>Backend Duration (sec)</b>: {row['TotalBackendCallDuration']}<br><b>Backend Calls:</b> <br>" + \
#                                      row['backendCallList'].replace(',', ',<br>')
#             else:
#                 concatenated_calls = f"<b>Backend Duration (sec)</b>: {row['TotalBackendCallDuration']}<br><b>Backend Calls:</b> None"
#
#             comparison_icon = ''
#             comparison_text = ''
#
#             if row['Compare'] == 'Higher':
#                 comparison_icon = '<i class="fas fa-arrow-up" title="Higher"></i>'
#                 comparison_text = 'Higher'
#             elif row['Compare'] == 'Lower':
#                 comparison_icon = '<i class="fas fa-arrow-down" title="Lower"></i>'
#                 comparison_text = 'Lower'
#             elif row['Compare'] == 'Equal':
#                 comparison_icon = '<i class="fas fa-exchange-alt" title="Equal"></i>'
#                 comparison_text = 'Equal'
#
#             data.append({
#                 'ApplicationName': row['ApplicationName'],
#                 'releaseID': row['releaseID'],
#                 'transactionName': row['transactionName'],
#                 'SLA': row['SLA'],
#                 'backendCall': concatenated_calls,
#                 'Compare': row['Compare']
#             })
#         return jsonify(data)
#
#     else:
#         # Handle GET request logic here (if needed)
#         conn = get_db_connection()
#         apps = conn.execute('SELECT DISTINCT applicationName FROM NFRDetails').fetchall()
#         conn.close()
#
#         searches_cookie = request.cookies.get('recent_searches')
#         if searches_cookie:
#             recent_searches = json.loads(searches_cookie)
#         else:
#             recent_searches = []
#
#         return render_template('discrepancy.html', apps=apps, recent_searches=recent_searches)

@app.route('/discrepancy', methods=['GET', 'POST'])
def handle_discrepancy():
    try:
        if request.method == 'POST':
            application_name = request.form.get('application_name')
            release_id = request.form.get('release_id')
            transaction_name = request.form.get('transaction_name')

            # If transaction_name is empty, set it to None
            transaction_name = transaction_name if transaction_name else None

            query = """
                            WITH LatestReleaseForBackendCall AS (
                                SELECT 
                                    transactionName AS backendCall, 
                                    MAX(releaseID) AS recentReleaseID
                                FROM NFRDetails
                                GROUP BY transactionName
                            ),
                            NFRDetailDepend AS (
                                SELECT
                                    t.applicationName, 
                                    t.transactionName, 
                                    t.releaseID, 
                                    t.SLA,
                                    d.backendCall, 
                                    d.callType
                                FROM NFRDetails t
                                LEFT JOIN NFROperationDependency d ON t.transactionName = d.transactionName
                                LEFT JOIN LatestReleaseForBackendCall lrb ON d.backendCall = lrb.backendCall
                                WHERE t.applicationName = ? 
                                  AND (t.releaseID = ? OR EXISTS (
                                        SELECT 1
                                        FROM LatestReleaseForBackendCall lrb
                                        WHERE lrb.backendCall = d.backendCall
                                          AND lrb.recentReleaseID = t.releaseID
                                    ))
                                  AND (? IS NULL OR t.transactionName LIKE ?)
                            ),
                            SLAComparisonCTE AS (
                                SELECT 
                                    t.ApplicationName, 
                                    t.transactionName, 
                                    t.backendCall, 
                                    t.CallType, 
                                    t.SLA, 
                                    t.releaseID,
                                    CASE 
                                        WHEN t.CallType = 'Async' THEN 
                                            (SELECT MAX(SLA) 
                                             FROM NFRDetails nr 
                                             WHERE nr.transactionName = t.backendCall 
                                               AND nr.releaseID = (SELECT MAX(releaseID) FROM NFRDetails WHERE transactionName = t.backendCall))
                                        WHEN t.CallType = 'Sync' THEN 
                                            (SELECT SUM(SLA) 
                                             FROM NFRDetails nr 
                                             WHERE nr.transactionName = t.backendCall 
                                               AND nr.releaseID = (SELECT MAX(releaseID) FROM NFRDetails WHERE transactionName = t.backendCall))
                                        ELSE 0 
                                    END AS SLAComparison
                                FROM NFRDetailDepend t
                            )
                            SELECT
                                ApplicationName, 
                                releaseID, 
                                transactionName, 
                                SLA,
                                ROUND(COALESCE(TotalSyncSLA, 0) + COALESCE(MaxAsyncSLA, 0), 2) AS TotalBackendCallDuration,
                                GROUP_CONCAT(backendCall) AS backendCallList,
                                CASE
                                    WHEN COALESCE(TotalSyncSLA, 0) + COALESCE(MaxAsyncSLA, 0) = 0 THEN 'NA'
                                    ELSE CASE 
                                        WHEN SLA > COALESCE(TotalSyncSLA, 0) + COALESCE(MaxAsyncSLA, 0) THEN 'Higher' 
                                        WHEN SLA < COALESCE(TotalSyncSLA, 0) + COALESCE(MaxAsyncSLA, 0) THEN 'Lower' 
                                        ELSE 'Equal' 
                                    END 
                                END AS Compare
                            FROM (
                                SELECT 
                                    ApplicationName, 
                                    transactionName, 
                                    releaseID, 
                                    SLA, 
                                    backendCall,
                                    SUM(CASE WHEN CallType = 'Sync' THEN SLAComparison ELSE 0 END) OVER (PARTITION BY transactionName, releaseID) AS TotalSyncSLA,
                                    MAX(CASE WHEN CallType = 'Async' THEN SLAComparison ELSE 0 END) OVER (PARTITION BY transactionName, releaseID) AS MaxAsyncSLA
                                FROM SLAComparisonCTE
                            ) AS p
                            WHERE ApplicationName = ? 
                              AND releaseID = ? 
                              AND (? IS NULL OR transactionName LIKE ?)
                            GROUP BY ApplicationName, releaseID, transactionName, SLA;
                        """

            params = [
                application_name,
                release_id,
                transaction_name, f'%{transaction_name}%' if transaction_name else None,
                application_name,
                release_id,
                transaction_name, f'%{transaction_name}%' if transaction_name else None
            ]

            conn = get_db_connection()
            cur = conn.cursor()
            cur.execute(query, params)
            rows = cur.fetchall()
            conn.close()

            data = []

            for row in rows:
                if row['backendCallList']:
                    concatenated_calls = f"<b>Backend Duration (sec)</b>: {row['TotalBackendCallDuration']}<br><b>Backend Calls:</b> <br>" + \
                                         row['backendCallList'].replace(',', ',<br>')
                else:
                    concatenated_calls = f"<b>Backend Duration (sec)</b>: {row['TotalBackendCallDuration']}<br><b>Backend Calls:</b> None"

                comparison_icon = ''
                comparison_text = ''

                if row['Compare'] == 'Higher':
                    comparison_icon = '<i class="fas fa-arrow-up" title="Higher"></i>'
                    comparison_text = 'Higher'
                elif row['Compare'] == 'Lower':
                    comparison_icon = '<i class="fas fa-arrow-down" title="Lower"></i>'
                    comparison_text = 'Lower'
                elif row['Compare'] == 'Equal':
                    comparison_icon = '<i class="fas fa-exchange-alt" title="Equal"></i>'
                    comparison_text = 'Equal'

                data.append({
                    'ApplicationName': row['ApplicationName'],
                    'releaseID': row['releaseID'],
                    'transactionName': row['transactionName'],
                    'SLA': row['SLA'],
                    'backendCall': concatenated_calls,
                    'Compare': row['Compare']
                })


            return jsonify(data)

        else:
            conn = get_db_connection()
            apps = conn.execute('SELECT DISTINCT applicationName FROM NFRDetails').fetchall()
            conn.close()

            searches_cookie = request.cookies.get('recent_searches')
            recent_searches = json.loads(searches_cookie) if searches_cookie else []

            return render_template('discrepancy.html', apps=apps, recent_searches=recent_searches)

    except Exception as e:
        print("Exception occurred:", str(e))
        return jsonify(success=False, message=str(e)), 500


# Save search handler with dynamic cookie name
@app.route('/save_search/<cookie_name>', methods=['POST'])
def save_search(cookie_name):
    data = request.json
    # print(data)

    # Retrieve the current searches from the cookie
    searches_cookie = request.cookies.get(cookie_name)
    if searches_cookie:
        recent_searches = json.loads(searches_cookie)
    else:
        recent_searches = []

    # Check if the search criteria already exists
    if data not in recent_searches:
        # Add new search criteria to the list
        recent_searches.append(data)
        # Keep only the last 5 searches
        recent_searches = recent_searches[-5:]

        # Set the cookie
        response = make_response(jsonify({'status': 'success'}))
        expires = datetime.utcnow() + timedelta(days=3)
        response.set_cookie(cookie_name, json.dumps(recent_searches), expires=expires)
    else:
        # If search criteria already exists, return success without updating the cookie
        response = jsonify({'status': 'success'})

    return response


# Route to get recent searches
@app.route('/recent_searches/<cookie_name>', methods=['GET'])
def recent_searches(cookie_name):
    recent_searches_cookie = request.cookies.get(cookie_name)
    if recent_searches_cookie:
        recent_searches = json.loads(recent_searches_cookie)
    else:
        recent_searches = []
    # print(recent_searches)
    # print(cookie_name)
    return jsonify(recent_searches)


# Function to process Excel file and update database

def process_excel(filepath):
    error_log = []
    total_records = 0
    inserted_records = 0
    updated_records = 0
    passed_records = 0
    failed_records = 0

    # Define mandatory columns
    mandatory_columns = ['ApplicationName', 'ReleaseID', 'BusinessScenario', 'TransactionName', 'BackendCall', 'CallType']

    try:
        wb = load_workbook(filename=filepath)
        ws = wb.active
    except Exception as load_error:
        error_log.append(f"Error loading Excel file: {load_error}")
        app.logger.error(f"Error loading Excel file: {load_error}", extra={'func': 'PROCESS_EXCEL_FILE'})
        return error_log, None  # Return error_log and summary as None immediately if loading the file fails

    # Extract headers from the first row
    headers = [cell.value for cell in ws[1]]

    # Check if mandatory columns are present
    missing_columns = [col for col in mandatory_columns if col not in headers]
    if missing_columns:
        error_message = f"Missing mandatory columns: {', '.join(missing_columns)}"
        error_log.append(error_message)
        app.logger.error(error_message, extra={'func': 'PROCESS_EXCEL_FILE'})
        raise Exception(error_message)  # Abort processing if mandatory columns are missing

    # Map headers to their indexes
    header_index_map = {header: index for index, header in enumerate(headers)}

    current_user = session.get('user', 'defaultUser')  # Retrieve current user
    current_time = "2024-06-06 10:10:10"  # datetime.now().strftime('%Y-%m-%d %H:%M:%S')

    with get_db_connection() as conn:
        for row_number, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            # Extract data using header indices
            try:
                applicationName = row[header_index_map['ApplicationName']]
                releaseID = row[header_index_map['ReleaseID']]
                businessScenario = row[header_index_map['BusinessScenario']]
                transactionName = row[header_index_map['TransactionName']]
                callType = row[header_index_map['CallType']]
                backendCall = row[header_index_map['BackendCall']]
                comments = row[header_index_map['Comments']] if 'Comments' in header_index_map else "" #None
            except IndexError as e:
                failed_records += 1
                error_message = f"Error processing row {row_number}: {e}"
                error_log.append(error_message)
                app.logger.error(error_message, extra={'func': 'PROCESS_EXCEL_FILE'})
                continue
            createdBy = "defaultUser"
            total_records += 1

            try:
                # Check if the record exists
                cursor = conn.execute('''
                    SELECT * FROM NFROperationDependency
                    WHERE applicationName = ? AND releaseID = ? AND  transactionName = ?
                ''', (applicationName, releaseID, transactionName))

                existing_record = cursor.fetchone()
                if existing_record:
                    # Update the record
                    conn.execute("""
                        UPDATE NFROperationDependency
						SET backendCall = ?, callType = ?, comments = ?,  modifiedBy = ?, modifed_date = ?
						WHERE applicationName = ? AND businessScenario = ? AND transactionName = ? AND releaseID = ?
						""", (backendCall, callType, comments,  createdBy, datetime.now(),
							  applicationName, businessScenario, transactionName, releaseID))
                    updated_records += 1

                else:
                    # Insert the record
                    conn.execute("""
                        INSERT INTO NFROperationDependency (applicationName, releaseID, businessScenario, transactionName, backendCall, callType, comments, createdBy, created_date)
					VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
					""", (applicationName, releaseID, businessScenario, transactionName, backendCall, callType, comments, createdBy, datetime.now()))
                    inserted_records += 1
                passed_records += 1
            except Exception as e:
                failed_records += 1
                error_message = f"Error processing row {row_number}: {e}"
                error_log.append(error_message)
                app.logger.error(error_message, extra={'func': 'PROCESS_EXCEL_FILE'})

        conn.commit()

    summary = {
        'total': total_records,
        'inserted': inserted_records,
        'updated': updated_records,
        'passed': passed_records,
        'failed': failed_records
    }
    print(summary)
    return error_log, summary



@app.route('/upload_dependency', methods=['GET','POST'])
def upload_dependency():
    error_log = []  # Initialize error_log
    summary = None  # Initialize summary
    if request.method == 'GET':
        pass
    if request.method == 'POST':
        file = request.files['file']
        if file and file.filename.endswith('.xlsx'):
            filepath = os.path.join(app.config['UPLOAD_FOLDER'], file.filename)
            file.save(filepath)
            try:
                error_log, summary = process_excel(filepath)
                if summary is None:
                    flash('File was not processed due to errors.', 'danger')
                else:
                    flash('File successfully uploaded and processed.', 'success')
            except Exception as e:
                app.logger.error(f'Error processing file: {e}', extra={'func': 'UPLOAD'})
                flash(f'Error processing file: {e}', 'danger')
            return render_template('upload_dependency.html', logs=error_log, summary=summary)
    return render_template('upload_dependency.html', logs=[], summary=summary)

# Function to check user role
def check_user_role(userid):
    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute('SELECT role FROM UserAccess WHERE userid = ?', (userid,))
    role = cursor.fetchone()
    conn.close()
    return role[0] if role else None

# Route to render the webpage for managing UserAccess
@app.route('/manage_user_access', methods=['GET'])
def manage_user_access():
    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute('SELECT * FROM UserAccess')
    users = cursor.fetchall()
    conn.close()
    return render_template('manage_user_access.html', users=users)

# Route to handle adding or editing UserAccess records
@app.route('/save_user_access', methods=['POST'])
def save_user_access():
    userid = request.form['userid']
    role = request.form['role']
    status = request.form['status']

    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute('INSERT OR REPLACE INTO UserAccess (userid, role, status) VALUES (?, ?, ?)', (userid, role, status))
    conn.commit()
    conn.close()

    return jsonify({'success': True})

# Route to handle deleting UserAccess records
@app.route('/delete_user_access/<userid>', methods=['DELETE'])
def delete_user_access(userid):
    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute('DELETE FROM UserAccess WHERE userid = ?', (userid,))
    conn.commit()
    conn.close()

    return jsonify({'success': True})


# User credentials for demo purposes (should be securely managed in a real application)
USER_CREDENTIALS = {
    'admin': 'admin123',
    'poweruser': 'power123',
    'user': 'user123'
}

# Function to check if credentials are valid
def validate_credentials(username, password):
    if username in USER_CREDENTIALS and USER_CREDENTIALS[username] == password:
        return True
    return False


# Route to handle login and authentication
@app.route('/login', methods=['GET','POST'])
def login():
    if request.method == 'GET':
        return render_template('login.html')
    if request.method == 'POST':
        username = request.form['username']
        password = request.form['password']

        if validate_credentials(username, password):
            role = check_user_role(username)
            return redirect(url_for('manage_user_access', role=role))
        else:
            return render_template('login.html', error='Invalid credentials. Please try again.')


if __name__ == '__main__':
    app.run(debug=True)
